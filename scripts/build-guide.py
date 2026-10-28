"""Build the downloadable Word guide from the same Markdown and prompt as GitHub.

Maintainers: use python-docx 1.x. After building, refresh the TOC in Word, save,
then run this script with --finalize to refresh document metadata and hashes.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import posixpath
import re
import tempfile
from urllib.parse import quote
from zipfile import ZipFile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/beginner-guide.md'
PROMPT = ROOT / 'prompts/带我完成配置.txt'
OUTPUT = ROOT / 'docs/downloads/防封指南_保姆联动版.docx'
MANIFEST = ROOT / 'docs/guide-version.json'
REPO = 'https://github.com/clm031102-source/claude-network-guide'
EDITION = '2026.10.28-r2'
BEGIN = '<!-- AGENT_PROMPT_START -->'
END = '<!-- AGENT_PROMPT_END -->'


def build_fingerprint() -> str:
    """Digest the inputs actually used to build the downloadable document."""
    inputs = [SOURCE, PROMPT]
    for relative in re.findall(r'!\[[^\]]*\]\(([^)]+)\)', SOURCE.read_text(encoding='utf-8')):
        inputs.append((SOURCE.parent / relative).resolve())
    digests = {path.relative_to(ROOT).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs}
    return hashlib.sha256(json.dumps(digests, sort_keys=True, ensure_ascii=False).encode('utf-8')).hexdigest()


def sync_prompt() -> str:
    prompt = PROMPT.read_text(encoding='utf-8').strip()
    source = SOURCE.read_text(encoding='utf-8')
    block = f'{BEGIN}\n\n```text\n{prompt}\n```\n\n{END}'
    if '<!-- AGENT_PROMPT_INSERT -->' in source:
        source = source.replace('<!-- AGENT_PROMPT_INSERT -->', block)
    elif BEGIN in source and END in source:
        source = re.sub(re.escape(BEGIN) + r'.*?' + re.escape(END), lambda _: block, source, flags=re.S)
    else:
        raise ValueError('The guide must contain a prompt synchronization marker.')
    # Stable chapter anchors survive heading wording changes in future revisions.
    source = re.sub(r'<a id="chapter-\d+"></a>\n\n', '', source)
    chapters = list(re.finditer(r'^## (.+)$', source, re.M))
    chapter_index = '\n'.join(f'{number}. [{match.group(1)}](#chapter-{number:02d})' for number, match in enumerate(chapters, 1))
    index_block = '<!-- CHAPTER_INDEX_START -->\n\n' + chapter_index + '\n\n<!-- CHAPTER_INDEX_END -->'
    if '<!-- CHAPTER_INDEX_START -->' in source:
        source = re.sub(r'<!-- CHAPTER_INDEX_START -->.*?<!-- CHAPTER_INDEX_END -->', lambda _: index_block, source, flags=re.S)
    else:
        source = source.replace('## 第一章', index_block + '\n\n## 第一章', 1)
    counter = 0
    def chapter_anchor(match):
        nonlocal counter
        counter += 1
        return f'<a id="chapter-{counter:02d}"></a>\n\n{match.group(0)}'
    source = re.sub(r'^## .+$', chapter_anchor, source, flags=re.M)
    SOURCE.write_text(source.rstrip() + '\n', encoding='utf-8', newline='\n')
    handoff_path = ROOT / 'docs/agent-handoff.md'
    handoff = handoff_path.read_text(encoding='utf-8')
    handoff, replacements = re.subn(r'```text\n.*?\n```', lambda _: f'```text\n{prompt}\n```', handoff, count=1, flags=re.S)
    if replacements != 1:
        raise ValueError('The handoff page must contain its primary text prompt block.')
    handoff_path.write_text(handoff.rstrip() + '\n', encoding='utf-8', newline='\n')
    return source


def online_url(value: str) -> str:
    if value.startswith(('https://', 'http://', 'mailto:')):
        return value
    path, sep, anchor = value.partition('#')
    relative = posixpath.normpath('docs/' + path) if path else 'docs/beginner-guide.md'
    return REPO + '/blob/main/' + quote(relative, safe='/') + ('#' + quote(anchor, safe='-') if sep else '')


def build() -> None:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    source = sync_prompt()
    document = Document()
    section = document.sections[0]
    # Keep the A4 geometry of the user's existing formatted guide.
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.85)
    section.different_first_page_header_footer = True

    def set_font(font, size=11, bold=False, mono=False, color='000000'):
        font.name = 'Consolas' if mono else 'Calibri'
        font.size, font.bold, font.italic = Pt(size), bold, False
        font.color.rgb = RGBColor.from_string(color)
        fonts = font._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), '微软雅黑')
        for name in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + name), None)

    specs = [('Normal', 11, False), ('Title', 29, True), ('Subtitle', 16, False),
             ('Heading 1', 18, True), ('Heading 2', 13, True), ('Heading 3', 11.5, True),
             ('TOC Heading', 22, True), ('TOC 1', 11, False), ('Code Guide', 9.3, False),
             ('Prompt Guide', 10.5, False), ('Agent Prompt', 10.5, False), ('Guide Caption', 9, False)]
    for name, size, bold in specs:
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(style.font, size, bold, name == 'Code Guide')
        pf = style.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(4)
        pf.line_spacing = Pt(17.5)
        pf.keep_with_next, pf.keep_together, pf.widow_control = False, False, True
        pr = style._element.get_or_add_pPr()
        for tag in ['w:pBdr', 'w:shd']:
            for el in list(pr.findall(qn(tag))):
                pr.remove(el)
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), 'false'); pr.append(snap)
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[name]
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != 'Heading 1' else 0)
        style.paragraph_format.space_after = Pt(6 if name != 'Heading 1' else 16)
        style.paragraph_format.line_spacing = Pt(25 if name == 'Heading 1' else 20)
    document.styles['Heading 1'].paragraph_format.page_break_before = True
    document.styles['Title'].paragraph_format.line_spacing = Pt(40)
    document.styles['Title'].paragraph_format.space_after = Pt(12)
    document.styles['Subtitle'].paragraph_format.line_spacing = Pt(25)
    document.styles['Subtitle'].paragraph_format.space_after = Pt(20)
    document.styles['TOC Heading'].paragraph_format.page_break_before = True
    document.styles['TOC Heading'].paragraph_format.space_after = Pt(20)
    toc_style = document.styles['TOC 1'].paragraph_format
    toc_style.space_after = Pt(10)
    toc_style.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    for name in ['Code Guide', 'Prompt Guide', 'Agent Prompt']:
        pf = document.styles[name].paragraph_format
        pf.left_indent = Cm(.25)
        pf.space_after = Pt(3)
        pf.line_spacing = Pt(14 if name == 'Code Guide' else 17)

    def hyperlink(paragraph, label, url):
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(online_url(url), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '185B75'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text)
        link.append(run); paragraph._p.append(link)

    pattern = re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`')
    def inline(paragraph, text):
        cursor = 0
        for match in pattern.finditer(text):
            if match.start() > cursor: paragraph.add_run(text[cursor:match.start()])
            if match.group(1) is not None: hyperlink(paragraph, match.group(1), match.group(2))
            elif match.group(3) is not None: paragraph.add_run(match.group(3)).bold = True
            else: paragraph.add_run(match.group(4)).font.name = 'Consolas'
            cursor = match.end()
        if cursor < len(text): paragraph.add_run(text[cursor:])

    def field(paragraph, instruction):
        run = OxmlElement('w:r'); begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); run.append(begin); paragraph._p.append(run)
        run = OxmlElement('w:r'); text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction; run.append(text); paragraph._p.append(run)
        run = OxmlElement('w:r'); end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); run.append(end); paragraph._p.append(run)

    header = section.header.paragraphs[0]
    header.add_run('防封指南  |  Claude Network Guide')
    set_font(header.runs[0].font, 8.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 '); field(footer, ' PAGE '); footer.add_run(' 页  /  共 '); field(footer, ' NUMPAGES '); footer.add_run(' 页')
    for run in footer.runs: set_font(run.font, 9)

    p = document.add_paragraph('防封指南', 'Title'); p.paragraph_format.space_before = Pt(45)
    document.add_paragraph('从购买服务器到 Claude 启动前检测', 'Subtitle')
    document.add_paragraph('零基础保姆教程  GitHub 联动版')
    document.add_paragraph(f'版本 {EDITION}  |  更新日期 2026 年 10 月 28 日')
    document.add_paragraph('按顺序完成采购、配置、检测和验收。每章都说明要做什么、如何判断完成，以及卡住后怎么继续。你也可以把最后一章的完整提示词交给信任的 AI 助手，让它一步一步带你做。')
    document.add_paragraph('这份指南用于减少配置错误和核对实际网络路径，不承诺账号不会受到限制。住宅 IP 是本项目选择的出口方案，不代替平台的地区和账号资格要求。')
    p = document.add_paragraph(); inline(p, '[在线最新版与完整教程](beginner-guide.md)  ·  [项目首页](' + REPO + ')  ·  [AI 接手提示词](agent-handoff.md)')
    document.add_paragraph(REPO, 'Guide Caption')
    document.add_paragraph('如何使用链接  Word 中按住 Ctrl 再单击链接即可打开对应资料；手机阅读时直接点按。目录可以跳转到本文件的章节，正文的在线入口可以回到仓库查最新步骤。', 'Normal')
    document.add_paragraph('目录与在线索引', 'TOC Heading')
    document.add_paragraph('先按目录找到当前步骤。首次配置建议从第一章开始；已有资源先核验再复用。', 'Normal')
    field(document.add_paragraph(), ' TOC \\o "1-1" \\h \\z ')

    def add_table(rows):
        table = document.add_table(rows=1, cols=len(rows[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = {2: [4.4, 12.2], 3: [3.2, 6.2, 7.2], 4: [2.5, 4.2, 4.8, 5.1]}.get(len(rows[0]), [16.6 / len(rows[0])] * len(rows[0]))
        for column, width in zip(table.columns, widths): column.width = Cm(width)
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for side, amount in [('top', 95), ('bottom', 95), ('left', 105), ('right', 105)]:
            el = OxmlElement('w:' + side); el.set(qn('w:w'), str(amount)); el.set(qn('w:type'), 'dxa'); margins.append(el)
        props.append(margins)
        for index, cells in enumerate(rows):
            row = table.rows[0] if index == 0 else table.add_row()
            pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
            if index == 0: pr.append(OxmlElement('w:tblHeader'))
            for n, cell in enumerate(row.cells):
                cell.width = Cm(widths[n]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF0F3' if index == 0 else ('F7F8F9' if index % 2 == 0 else 'FFFFFF')); cell._tc.get_or_add_tcPr().append(shade)
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(15)
                inline(p, cells[n] if n < len(cells) else '')
                for run in p.runs: set_font(run.font, 10, index == 0)
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    lines = source.splitlines(); index = 0; in_prompt = False; started = False; chapter_number = 0
    while index < len(lines):
        text = lines[index].strip(); index += 1
        if not text: continue
        if text.startswith('## '): started = True
        if not started or text.startswith('<a id='): continue
        if text == BEGIN:
            in_prompt = True
            p = document.add_paragraph('完整提示词  复制以下全部文字', 'Heading 2')
            p.paragraph_format.page_break_before = True
            continue
        if text == END: in_prompt = False; continue
        if text.startswith('<!--') or text == '---': continue
        if text.startswith('# '): continue
        if text.startswith('```'):
            style = 'Agent Prompt' if in_prompt else ('Prompt Guide' if text[3:].strip() == 'text' else 'Code Guide')
            while index < len(lines) and not lines[index].strip().startswith('```'):
                line = lines[index]; index += 1
                if line.strip(): document.add_paragraph(line, style)
                elif style in ['Prompt Guide', 'Agent Prompt']: document.add_paragraph().paragraph_format.space_after = Pt(1)
            index += 1; continue
        if text.startswith('|') and index < len(lines) and re.match(r'^\|?[\s:|-]+\|?$', lines[index].strip()):
            rows = [[c.strip() for c in text.strip('|').split('|')]]; index += 1
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append([c.strip() for c in lines[index].strip().strip('|').split('|')]); index += 1
            add_table(rows); continue
        image_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', text)
        if image_match:
            image = (SOURCE.parent / image_match.group(2)).resolve()
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1
            p.add_run().add_picture(str(image), width=Cm(16.6))
            document.add_paragraph(image_match.group(1), 'Guide Caption'); continue
        heading = re.match(r'^(#{2,4})\s+(.+)$', text)
        if heading:
            level = len(heading.group(1)) - 1
            name = re.sub(r'[：:；;，,。！？!?（）()·—/、“”‘’]', ' ', heading.group(2))
            p = document.add_paragraph(name, 'Heading ' + str(level))
            if heading.group(2) in ['DNS、UDP 与 IPv6 要分别检查', '根据需求选择订阅，不按“养号日程”付款', '维护时保留旧版本', '写一份别人看得懂的记录']:
                p.paragraph_format.page_break_before = True
            if level == 1:
                chapter_number += 1
                p = document.add_paragraph(style='Guide Caption')
                hyperlink(p, '本章在线版  查看最新步骤', f'beginner-guide.md#chapter-{chapter_number:02d}')
                p.paragraph_format.keep_with_next = True
            continue
        p = document.add_paragraph()
        if text.startswith('> '): text = text[2:]
        item = re.match(r'^(?:[-*]\s+|\d+[.)、]\s+)(.*)$', text)
        if item:
            marker = '• ' if text[0] in '-*' else text[:len(text)-len(item.group(1))]
            p.paragraph_format.left_indent = Cm(.45); p.paragraph_format.first_line_indent = Cm(-.45)
            inline(p, marker + item.group(1))
        else: inline(p, text)

    document.core_properties.title = '防封指南 从购买服务器到 Claude 启动前检测'
    document.core_properties.subject = '零基础教程 与 GitHub 仓库联动'
    document.core_properties.author = 'clm031102-source'
    document.core_properties.last_modified_by = 'clm031102-source'
    document.core_properties.keywords = 'claude-network-guide-build-v1:' + build_fingerprint()
    document.core_properties.version = EDITION
    document.core_properties.comments = 'Generated from docs/beginner-guide.md and prompts/带我完成配置.txt'
    settings = document.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f'Built {OUTPUT.name}; refresh the table of contents in Word before publishing.')


def finalize() -> None:
    # Word may set the current Windows account as lastModifiedBy; use repo attribution.
    namespaces = {'dc': 'http://purl.org/dc/elements/1.1/', 'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'}
    # Keep the dcterms prefix because xsi:type values refer to it by name.
    for prefix, uri in {**namespaces, 'dcterms': 'http://purl.org/dc/terms/',
                        'xsi': 'http://www.w3.org/2001/XMLSchema-instance'}.items():
        ET.register_namespace(prefix, uri)
    with ZipFile(OUTPUT) as src:
        members = [(info, src.read(info.filename)) for info in src.infolist()]
        core = ET.fromstring(src.read('docProps/core.xml'))
        stamp = core.find('cp:keywords', namespaces)
        if stamp is None or stamp.text != 'claude-network-guide-build-v1:' + build_fingerprint():
            raise ValueError('Word was not built from the current guide, prompt and images. Run the builder again before finalizing.')
    descriptor, name = tempfile.mkstemp(suffix='.docx', dir=OUTPUT.parent)
    import os
    os.close(descriptor)
    try:
        with ZipFile(name, 'w') as target:
            for info, content in members:
                if info.filename == 'docProps/core.xml':
                    xml = ET.fromstring(content)
                    for tag in ['dc:creator', 'cp:lastModifiedBy']:
                        element = xml.find(tag, namespaces)
                        if element is not None: element.text = 'clm031102-source'
                    content = ET.tostring(xml, encoding='utf-8', xml_declaration=True)
                target.writestr(info, content)
        os.replace(name, OUTPUT)
    finally:
        Path(name).unlink(missing_ok=True)
    manifest = {'edition': EDITION, 'updated': '2026-10-28', 'repository': REPO,
                'source': SOURCE.relative_to(ROOT).as_posix(), 'prompt': PROMPT.relative_to(ROOT).as_posix(),
                'word': OUTPUT.relative_to(ROOT).as_posix(), 'sha256': {}}
    for path in [SOURCE, PROMPT, OUTPUT]:
        manifest['sha256'][path.relative_to(ROOT).as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    print('Guide version and source/prompt/Word hashes recorded.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--finalize', action='store_true', help='Only sanitize metadata and update release hashes after Word saves the TOC')
    args = parser.parse_args()
    if not args.finalize: build()
    finalize()
